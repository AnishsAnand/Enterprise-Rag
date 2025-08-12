from fastapi import APIRouter, HTTPException, BackgroundTasks, UploadFile, File
from pydantic import BaseModel, HttpUrl
from typing import List, Dict, Any
import asyncio
from datetime import datetime
import mimetypes
from io import BytesIO, StringIO
import csv
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
from bs4 import BeautifulSoup

from app.services.scraper_service import scraper_service
from app.services.chroma_service import chroma_service
from app.services.ai_service import ai_service

router = APIRouter()



class WidgetQueryRequest(BaseModel):
    query: str
    max_results: int = 15  
    include_sources: bool = True

class WidgetScrapeRequest(BaseModel):
    url: HttpUrl
    store_in_knowledge: bool = True

class BulkScrapeRequest(BaseModel):
    base_url: HttpUrl
    max_depth: int = 10    
    max_urls: int = 300    
    auto_store: bool = True

def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 200) -> List[str]:
    chunks = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks


@router.post("/widget/query")
async def widget_query(request: WidgetQueryRequest):
    try:
        print(f"Widget query received: {request.query}")
        search_results = await chroma_service.search_documents(request.query, request.max_results)

        filtered_results = [r for r in search_results if r['relevance_score'] > 0.7]
        if not filtered_results:
            filtered_results = search_results[:3] 

        context = [r['content'] for r in filtered_results]

        answer = await ai_service.generate_response(request.query, context)

        sources = []
        if request.include_sources:
            for result in filtered_results:
                sources.append({
                    'url': result['metadata'].get('url', ''),
                    'title': result['metadata'].get('title', 'Untitled'),
                    'relevance_score': result['relevance_score'],
                    'content_preview': (result['content'][:200] + '...') if len(result['content']) > 200 else result['content']
                })

        return {
            'query': request.query,
            'answer': answer,
            'sources': sources,
            'has_sources': len(sources) > 0,
            'timestamp': datetime.now().isoformat()
        }

    except Exception as e:
        print(f"Widget query error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/widget/scrape")
async def widget_scrape(request: WidgetScrapeRequest, background_tasks: BackgroundTasks):
    try:
        print(f"Widget scrape request: {request.url}")

        scrape_params = {
            'extract_text': True,
            'extract_links': False,
            'extract_images': False,
            'extract_tables': True,
            'scroll_page': True
        }

        result = await scraper_service.scrape_url(str(request.url), scrape_params)

        if result['status'] == 'success':
            if request.store_in_knowledge:
                document = {
                    'content': result['content'].get('text', ''),
                    'url': str(request.url),
                    'title': result['content'].get('title', ''),
                    'format': 'text',
                    'timestamp': datetime.now().isoformat(),
                    'source': 'widget_scrape'
                }
                background_tasks.add_task(store_document_task, document)

            return {
                'status': 'success',
                'url': str(request.url),
                'title': result['content'].get('title', 'Untitled'),
                'content_length': len(result['content'].get('text', '')),
                'method_used': result['method'],
                'stored_in_knowledge': request.store_in_knowledge,
                'timestamp': result['timestamp']
            }
        else:
            raise HTTPException(status_code=400, detail=result.get('error', 'Scraping failed'))

    except Exception as e:
        print(f"Widget scrape error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/widget/bulk-scrape")
async def widget_bulk_scrape(request: BulkScrapeRequest, background_tasks: BackgroundTasks):
    try:
        print(f"Widget bulk scrape request: {request.base_url}")
        discovered_urls = await scraper_service.discover_urls(
            str(request.base_url),
            request.max_depth,
            request.max_urls
        )

        if not discovered_urls:
            return {
                'status': 'no_urls_found',
                'message': 'No URLs discovered from the base URL',
                'base_url': str(request.base_url)
            }

        background_tasks.add_task(bulk_scrape_widget_task, discovered_urls, request.auto_store)

        return {
            'status': 'started',
            'base_url': str(request.base_url),
            'discovered_urls_count': len(discovered_urls),
            'urls_preview': discovered_urls[:5],
            'auto_store': request.auto_store,
            'estimated_time': f"{len(discovered_urls) * 2} seconds"
        }

    except Exception as e:
        print(f"Widget bulk scrape error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/widget/upload-file")
async def widget_upload_file(file: UploadFile = File(...), store_in_knowledge: bool = True):
    try:
        filename = file.filename
        content_type = file.content_type or mimetypes.guess_type(filename)[0]

        content = await file.read()
        text = None
        format_type = None

        if content_type == "text/plain":
            text = content.decode("utf-8")
            format_type = "text"
        elif content_type == "application/pdf":
            pdf_reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in pdf_reader.pages)
            format_type = "pdf"
        elif content_type == "text/csv":
            s = StringIO(content.decode("utf-8"))
            reader = csv.reader(s)
            text = "\n".join([", ".join(row) for row in reader])
            format_type = "csv"
        elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            format_type = "docx"
        elif content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
            text = "\n".join(
                ", ".join(str(cell.value) for cell in row if cell.value is not None)
                for sheet in wb.worksheets
                for row in sheet.iter_rows()
            )
            format_type = "xlsx"
        elif content_type in ["text/html", "application/xhtml+xml"]:
            soup = BeautifulSoup(content, "html.parser")
            text = soup.get_text(separator="\n")
            format_type = "html"
        else:
            raise HTTPException(status_code=415, detail=f"Unsupported file type: {content_type}")

        if not text or len(text.strip()) == 0:
            raise HTTPException(status_code=400, detail="File is empty or unreadable.")

        response = {
            "filename": filename,
            "format": format_type,
            "content_length": len(text),
            "stored_in_knowledge": False
        }

        if store_in_knowledge:
            chunks = chunk_text(text)
            documents_to_store = []
            for i, chunk in enumerate(chunks):
                documents_to_store.append({
                    "content": chunk,
                    "url": f"file://{filename}#chunk{i}",
                    "title": filename,
                    "format": format_type,
                    "timestamp": datetime.now().isoformat(),
                    "source": "widget_upload"
                })
            await chroma_service.add_documents(documents_to_store)
            response["stored_in_knowledge"] = True

        return response

    except Exception as e:
        print(f"Widget upload error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/widget/knowledge-stats")
async def widget_knowledge_stats():
    try:
        stats = await chroma_service.get_collection_stats()
        return {
            'document_count': stats.get('document_count', 0),
            'status': stats.get('status', 'unknown'),
            'last_updated': datetime.now().isoformat()
        }
    except Exception as e:
        print(f"Widget stats error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.delete("/widget/clear-knowledge")
async def widget_clear_knowledge():
    try:
        await chroma_service.delete_collection()
        await chroma_service.initialize()
        return {
            'status': 'success',
            'message': 'Knowledge base cleared successfully',
            'timestamp': datetime.now().isoformat()
        }
    except Exception as e:
        print(f"Widget clear error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

async def store_document_task(document: Dict[str, Any]):
    try:
        await chroma_service.add_documents([document])
        print(f"Stored document in knowledge base: {document.get('url', '')}")
    except Exception as e:
        print(f"Error storing document: {e}")

async def bulk_scrape_widget_task(urls: List[str], auto_store: bool):
    scraped_count = 0
    stored_count = 0
    batch_size = 3

    for i in range(0, len(urls), batch_size):
        batch = urls[i:i + batch_size]

        scrape_params = {
            'extract_text': True,
            'extract_links': False,
            'extract_images': False,
            'extract_tables': True,
            'scroll_page': True
        }

        tasks = [scraper_service.scrape_url(url, scrape_params) for url in batch]
        results = await asyncio.gather(*tasks, return_exceptions=True)

        documents_to_store = []

        for url, result in zip(batch, results):
            if isinstance(result, Exception):
                print(f"Error scraping {url}: {str(result)}")
                continue

            if result.get('status') == 'success':
                scraped_count += 1

                if auto_store:
                    document = {
                        'content': result['content'].get('text', ''),
                        'url': url,
                        'title': result['content'].get('title', ''),
                        'format': 'text',
                        'timestamp': datetime.now().isoformat(),
                        'source': 'widget_bulk_scrape'
                    }
                    documents_to_store.append(document)

        if documents_to_store:
            try:
                await chroma_service.add_documents(documents_to_store)
                stored_count += len(documents_to_store)
                print(f"Stored {len(documents_to_store)} documents from batch")
            except Exception as e:
                print(f"Error storing batch documents: {e}")

        await asyncio.sleep(2) 

    print(f"Bulk scrape completed: {scraped_count} scraped, {stored_count} stored")
